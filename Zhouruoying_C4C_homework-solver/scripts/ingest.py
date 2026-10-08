"""阶段 ①：摄取。把 .md/.txt/.tex/.pdf/.docx 统一读成带来源锚点的纯文本。

不做任何语义解释 —— 解析层只拿到 normalized text + 行号映射，便于溯源。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

SUPPORTED_TEXT = {".md", ".txt", ".markdown", ".tex"}
SUPPORTED_PDF = {".pdf"}
SUPPORTED_DOCX = {".docx"}


def _lines_with_anchor(text: str) -> list[dict]:
    out = []
    for i, line in enumerate(text.splitlines(), start=1):
        out.append({"line": i, "text": line})
    return out


def _read_pdf(path: Path) -> tuple[str, dict]:
    meta = {"extractor": None, "pages": 0}
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError:
            raise ModuleNotFoundError(
                "读取 PDF 需要 pdfplumber 或 pypdf：pip install pdfplumber>=0.9")
        r = PdfReader(str(path))
        meta["extractor"] = "pypdf"
        meta["pages"] = len(r.pages)
        return "\n".join((pg.extract_text() or "") for pg in r.pages), meta
    parts = []
    with pdfplumber.open(str(path)) as pdf:
        meta["extractor"] = "pdfplumber"
        meta["pages"] = len(pdf.pages)
        for pg in pdf.pages:
            parts.append(pg.extract_text() or "")
    return "\n".join(parts), meta


def _read_docx(path: Path) -> tuple[str, dict]:
    try:
        import docx  # type: ignore
    except ImportError:
        raise ModuleNotFoundError(
            "读取 .docx 需要 python-docx：pip install python-docx>=0.8")
    d = docx.Document(str(path))
    paras = [p.text for p in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            paras.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(paras), {"extractor": "python-docx",
                              "paragraphs": len(paras)}


def _strip_tex_wrapper(text: str) -> str:
    """若是 .tex 输入（例如 starter 的 challenge_answers_cn.tex），
    剥掉导言区与 \\begin{document} 包裹，只留正文供解析。"""
    m = re.search(r"\\begin\{document\}(.*?)\\end\{document\}", text, re.S)
    body = m.group(1) if m else text
    body = re.sub(r"%.*", "", body)              # 去注释
    body = re.sub(r"\\(section|subsection)\*?\{([^}]*)\}",
                  r"\n## \2\n", body)            # 章节转 md 标题
    body = re.sub(r"\\item\s*", "\n- ", body)
    body = body.replace("\\\\", "\n")
    return body.strip()


def ingest(path_str: str) -> dict:
    path = Path(path_str).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"输入文件不存在：{path}")
    suffix = path.suffix.lower()
    meta: dict = {"path": str(path), "suffix": suffix,
                  "size_bytes": path.stat().st_size}
    if suffix in SUPPORTED_TEXT:
        raw = path.read_text(encoding="utf-8", errors="replace")
        fmt = "text"
        if suffix == ".tex":
            raw = _strip_tex_wrapper(raw)
            fmt = "tex"
        meta["format"] = fmt
    elif suffix in SUPPORTED_PDF:
        raw, pm = _read_pdf(path)
        meta.update(pm)
        meta["format"] = "pdf"
    elif suffix in SUPPORTED_DOCX:
        raw, dm = _read_docx(path)
        meta.update(dm)
        meta["format"] = "docx"
    else:
        raise ValueError(
            f"不支持的输入类型 {suffix}；支持 {sorted(SUPPORTED_TEXT | SUPPORTED_PDF | SUPPORTED_DOCX)}")

    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    lines = _lines_with_anchor(raw)
    return {
        "meta": meta,
        "text": raw,
        "lines": lines,
        "n_lines": len(lines),
        "n_chars": len(raw),
        "sha256": __import__("hashlib").sha256(
            raw.encode("utf-8")).hexdigest(),
    }


if __name__ == "__main__":  # 手工调试
    import sys
    print(json.dumps(ingest(sys.argv[1])["meta"], ensure_ascii=False, indent=2))
